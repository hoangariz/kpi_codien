import os
import sys

# Ensure UTF-8 output in Windows terminal
sys.stdout.reconfigure(encoding='utf-8')

import openpyxl
import pandas as pd

excel_path = "demofile.xlsx"
print("File exists:", os.path.exists(excel_path))
print("File size:", os.path.getsize(excel_path), "bytes")

wb = openpyxl.load_workbook(excel_path, read_only=True)
print("Sheets:", wb.sheetnames)
ws = wb.active

rows = ws.iter_rows(values_only=True)
header = next(rows)
print(f"\nHeader columns ({len(header)}):")
for i, col in enumerate(header):
    print(f"  [{i}] {col}")

# Read first 2 sample rows
print("\nSample rows:")
for r_idx in range(2):
    row = next(rows, None)
    if not row:
        break
    print(f"\n--- Row {r_idx + 1} ---")
    for i, (col, val) in enumerate(zip(header, row)):
        if val is not None:
            print(f"  {col}: {val}")

# Check with pandas
df = pd.read_excel(excel_path, engine="openpyxl")
print(f"\nTotal rows in demofile.xlsx: {len(df)}")
print("\nDataFrame columns:")
for c in df.columns:
    print(" -", repr(c))

if "Loại công việc" in df.columns:
    print("\n--- Value counts: Loại công việc ---")
    print(df["Loại công việc"].value_counts().head(10))

target_task_type = "Bảo dưỡng cứng cơ điện điều hòa, máy phát điện, thông gió lọc bụi ICMS"
matching = df[df["Loại công việc"] == target_task_type]
print(f"\nTarget task type: '{target_task_type}'")
print(f"Total matching rows: {len(matching)}")

if len(matching) > 0:
    print("\nMatching - Trạng thái breakdown:")
    print(matching["Trạng thái"].value_counts(dropna=False))
    print("\nMatching - Top 10 Nhân viên thực hiện:")
    print(matching["Nhân viên thực hiện"].value_counts(dropna=False).head(10))
    print("\nMatching - Top 10 Nhóm điều phối:")
    print(matching["Nhóm điều phối"].value_counts(dropna=False).head(10))

if "Hệ thống" in df.columns:
    print("\n--- Value counts: Hệ thống (SPM / SPM_VTNET) ---")
    print(df["Hệ thống"].value_counts(dropna=False).head(10))
