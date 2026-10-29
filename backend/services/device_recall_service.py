import io
import csv
import os
from datetime import datetime, timedelta
from typing import Dict, Any, List, Optional
from sqlalchemy import func, distinct
from sqlalchemy.orm import Session

from backend.models.device_recall import DeviceRecall
from backend.models.settings import SystemSetting
from backend.config import BASE_DIR


def _set_setting(db: Session, key: str, value: str):
    item = db.query(SystemSetting).filter(SystemSetting.key == key).first()
    if item:
        item.value = value
        item.updated_at = datetime.utcnow()
    else:
        db.add(SystemSetting(key=key, value=value))
    db.commit()


def _get_setting(db: Session, key: str, default: str = "") -> str:
    item = db.query(SystemSetting).filter(SystemSetting.key == key).first()
    return item.value if item else default


def extract_short_cluster(cum_xa: str) -> str:
    """Rút gọn mã/tên cụm xã, ví dụ: HTH-007-SGG -> SGG, Trung tâm Bắc Hồng Lĩnh -> Bắc Hồng Lĩnh."""
    if not cum_xa:
        return ""
    name = cum_xa.strip()
    if name.lower().startswith("trung tâm "):
        return name[10:].strip()
    parts = name.split("-")
    if len(parts) >= 3:
        return parts[-1].strip().upper()
    return name


def parse_device_recall_text(content: str) -> tuple[List[dict], int, int]:
    """
    Phân tích nội dung file .txt dạng TSV có sẵn header.
    Tự động hỗ trợ cả format mới (thuhoithietbi_update.txt) và format cũ (thuhoithietbi.txt).
    Loại bỏ các dòng không có 'Trung tâm Cụm xã' hoặc thiếu số thuê bao.
    Trả về: (records_to_insert, total_rows, skipped_no_cluster)
    """
    f = io.StringIO(content)
    reader = csv.reader(f, delimiter='\t')
    rows = list(reader)

    if not rows:
        return [], 0, 0

    header_row = rows[0]
    total_raw_rows = len(rows) - 1

    # Map column headers accurately with specific prioritization
    col_map = {
        "so_thue_bao": None,
        "dich_vu": None,
        "dia_chi_khach_hang": None,
        "dien_thoai_lh": None,
        "user_thanh_ly": None,
        "muc_thiet_bi": None,
        "tong_tbi_phai_thu": None,
        "sl_thiet_bi_da_thu": None,
        "loai_thiet_bi": None,
        "tinh": None,
        "cum_xa": None,
        "ten_ft": None,
        "ma_nv": None,
        # Legacy columns
        "muc_thue_bao": None,
        "tuoi_tho_mesh": None,
        "so_tbi_mesh_phai_thu": None,
    }

    for idx, col in enumerate(header_row):
        c = str(col).strip().lower().replace('\n', ' ')
        if "đã thu" in c or "da thu" in c:
            col_map["sl_thiet_bi_da_thu"] = idx
        elif "tuổi thọ" in c:
            col_map["tuoi_tho_mesh"] = idx
        elif "mesh phải thu" in c or "số tbi mesh" in c or "so tbi mesh" in c:
            col_map["so_tbi_mesh_phai_thu"] = idx
        elif "mức thiết bị" in c or "muc thiet bi" in c:
            col_map["muc_thiet_bi"] = idx
        elif "mức thuê bao" in c or "muc thue bao" in c:
            col_map["muc_thue_bao"] = idx
        elif "tổng tbi" in c or "tổng thiết bị" in c or ("phải thu" in c and "mesh" not in c and "đã thu" not in c):
            col_map["tong_tbi_phai_thu"] = idx
        elif "số thuê bao" in c or c in ("số tb", "so_thue_bao", "thuê bao", "so thue bao"):
            col_map["so_thue_bao"] = idx
        elif "điện thoại" in c or "dien thoai" in c or "liên hệ" in c or "lien he" in c or c == "sđt":
            col_map["dien_thoai_lh"] = idx
        elif "thanh lý" in c or "thanh ly" in c:
            col_map["user_thanh_ly"] = idx
        elif "thường-0" in c or "tb thường" in c or "loại thiết bị" in c or ("mesh" in c and "thường" in c):
            col_map["loai_thiet_bi"] = idx
        elif "tỉnh" in c or "tinh" in c:
            col_map["tinh"] = idx
        elif "trung tâm" in c or "cụm xã" in c or "cum xa" in c:
            col_map["cum_xa"] = idx
        elif "tên ft" in c or ("nhân viên" in c and "mã" not in c and "mnv" not in c) or c == "ft":
            col_map["ten_ft"] = idx
        elif "mnv" in c or "mã nv" in c or "ma nv" in c:
            col_map["ma_nv"] = idx
        elif "dịch vụ" in c or "dich_vu" in c or "dich vu" in c:
            col_map["dich_vu"] = idx
        elif "địa chỉ" in c or "dia chi" in c:
            col_map["dia_chi_khach_hang"] = idx

    # Fallback checks if header format differs slightly
    if col_map["so_thue_bao"] is None:
        col_map["so_thue_bao"] = 1 if len(header_row) > 1 else 0
    if col_map["cum_xa"] is None:
        for idx, col in enumerate(header_row):
            if "cụm" in str(col).lower() or "tâm" in str(col).lower():
                col_map["cum_xa"] = idx
                break

    records = []
    skipped_no_cluster = 0
    now = datetime.utcnow()

    for row in rows[1:]:
        if not row:
            continue

        idx_cum = col_map["cum_xa"]
        cum_xa = str(row[idx_cum]).strip() if (idx_cum is not None and len(row) > idx_cum) else ""
        # YÊU CẦU: Xoá / bỏ qua các dòng không có Trung tâm Cụm xã
        if not cum_xa:
            skipped_no_cluster += 1
            continue

        idx_stb = col_map["so_thue_bao"]
        stb = str(row[idx_stb]).strip() if (idx_stb is not None and len(row) > idx_stb) else ""
        if not stb:
            continue

        idx_dv = col_map["dich_vu"]
        dich_vu = str(row[idx_dv]).strip() if (idx_dv is not None and len(row) > idx_dv) else None

        idx_dc = col_map["dia_chi_khach_hang"]
        dia_chi = str(row[idx_dc]).strip() if (idx_dc is not None and len(row) > idx_dc) else None

        idx_dt = col_map["dien_thoai_lh"]
        dien_thoai = str(row[idx_dt]).strip() if (idx_dt is not None and len(row) > idx_dt) else None

        idx_utl = col_map["user_thanh_ly"]
        user_tl = str(row[idx_utl]).strip() if (idx_utl is not None and len(row) > idx_utl) else None

        idx_mtbi = col_map["muc_thiet_bi"]
        muc_tbi = str(row[idx_mtbi]).strip() if (idx_mtbi is not None and len(row) > idx_mtbi) else None

        idx_mtb = col_map["muc_thue_bao"]
        muc_tb = str(row[idx_mtb]).strip() if (idx_mtb is not None and len(row) > idx_mtb) else None

        # Tổng tbi phải thu (int)
        idx_ttbi = col_map["tong_tbi_phai_thu"]
        raw_ttbi = str(row[idx_ttbi]).strip() if (idx_ttbi is not None and len(row) > idx_ttbi) else "0"
        try:
            tong_tbi = int(float(raw_ttbi))
        except Exception:
            tong_tbi = 0

        # SL thiết bị đã thu (int)
        idx_da_thu = col_map["sl_thiet_bi_da_thu"]
        raw_da_thu = str(row[idx_da_thu]).strip() if (idx_da_thu is not None and len(row) > idx_da_thu) else "0"
        try:
            sl_da_thu = int(float(raw_da_thu))
        except Exception:
            sl_da_thu = 0

        idx_ltb = col_map["loai_thiet_bi"]
        loai_tb = str(row[idx_ltb]).strip() if (idx_ltb is not None and len(row) > idx_ltb) else None

        idx_ttm = col_map["tuoi_tho_mesh"]
        tuoi_tho = str(row[idx_ttm]).strip() if (idx_ttm is not None and len(row) > idx_ttm) else None

        idx_stbm = col_map["so_tbi_mesh_phai_thu"]
        stb_mesh = str(row[idx_stbm]).strip() if (idx_stbm is not None and len(row) > idx_stbm) else None

        idx_tinh = col_map["tinh"]
        tinh = str(row[idx_tinh]).strip() if (idx_tinh is not None and len(row) > idx_tinh) else None

        idx_ft = col_map["ten_ft"]
        ten_ft = str(row[idx_ft]).strip() if (idx_ft is not None and len(row) > idx_ft) else "Chưa gán FT"

        idx_mnv = col_map["ma_nv"]
        ma_nv = str(row[idx_mnv]).strip() if (idx_mnv is not None and len(row) > idx_mnv) else None

        records.append({
            "so_thue_bao": stb,
            "dich_vu": dich_vu or None,
            "dia_chi_khach_hang": dia_chi or None,
            "dien_thoai_lh": dien_thoai or None,
            "user_thanh_ly": user_tl or None,
            "muc_thue_bao": muc_tb or None,
            "muc_thiet_bi": muc_tbi or None,
            "tong_tbi_phai_thu": tong_tbi,
            "sl_thiet_bi_da_thu": sl_da_thu,
            "loai_thiet_bi": loai_tb or None,
            "tuoi_tho_mesh": tuoi_tho or None,
            "so_tbi_mesh_phai_thu": stb_mesh or None,
            "tinh": tinh or None,
            "cum_xa": cum_xa,
            "ten_ft": ten_ft or "Chưa gán FT",
            "ma_nv": ma_nv or None,
            "created_at": now,
        })

    return records, total_raw_rows, skipped_no_cluster


def import_device_recall_content(db: Session, content: str, filename: str) -> dict:
    """
    Import danh sách thiết bị cần thu hồi từ chuỗi nội dung văn bản.
    Xóa sạch dữ liệu cũ và chèn danh sách mới.
    """
    records, total_rows, skipped_no_cluster = parse_device_recall_text(content)

    if not records:
        raise ValueError(
            "Không tìm thấy dòng dữ liệu hợp lệ nào có chứa 'Trung tâm Cụm xã'. "
            "Vui lòng kiểm tra lại định dạng file!"
        )

    # Gán filename
    for r in records:
        r["import_filename"] = filename

    # Xóa sạch dữ liệu cũ
    db.query(DeviceRecall).delete(synchronize_session=False)
    db.commit()

    # Bulk insert theo batch 1000 dòng
    BATCH_SIZE = 1000
    for i in range(0, len(records), BATCH_SIZE):
        batch = records[i : i + BATCH_SIZE]
        db.bulk_insert_mappings(DeviceRecall, batch)
        db.commit()

    # Thống kê metadata
    unique_clusters = sorted(list(set(r["cum_xa"] for r in records)))
    unique_fts = sorted(list(set(r["ten_ft"] for r in records)))
    total_devices = sum(r["tong_tbi_phai_thu"] for r in records)
    total_collected = sum(r.get("sl_thiet_bi_da_thu", 0) for r in records)

    now = datetime.utcnow()
    vn_now = now + timedelta(hours=7)
    vn_time_str = vn_now.strftime("%d/%m/%Y lúc %H:%M")

    _set_setting(db, "device_recall_last_import_time", now.isoformat())
    _set_setting(db, "device_recall_last_import_time_vn", vn_time_str)
    _set_setting(db, "device_recall_last_filename", filename)
    _set_setting(db, "device_recall_total_count", str(len(records)))
    _set_setting(db, "device_recall_skipped_count", str(skipped_no_cluster))
    _set_setting(db, "device_recall_clusters_count", str(len(unique_clusters)))
    _set_setting(db, "device_recall_fts_count", str(len(unique_fts)))
    _set_setting(db, "device_recall_total_devices", str(total_devices))
    _set_setting(db, "device_recall_total_collected", str(total_collected))

    return {
        "filename": filename,
        "total_raw_rows": total_rows,
        "valid_rows": len(records),
        "skipped_no_cluster": skipped_no_cluster,
        "total_clusters": len(unique_clusters),
        "total_fts": len(unique_fts),
        "total_devices": total_devices,
        "total_collected": total_collected,
        "remaining_devices": max(0, total_devices - total_collected),
        "clusters": unique_clusters,
        "imported_at_vn": vn_time_str,
    }


def get_device_recall_clusters(db: Session) -> List[dict]:
    """
    Lấy danh sách các cụm UNIQUE trong Trung tâm Cụm xã cùng thống kê tổng quan.
    """
    rows = (
        db.query(
            DeviceRecall.cum_xa,
            func.count(DeviceRecall.id).label("subscribers_count"),
            func.sum(DeviceRecall.tong_tbi_phai_thu).label("total_devices"),
            func.sum(DeviceRecall.sl_thiet_bi_da_thu).label("total_collected"),
            func.count(distinct(DeviceRecall.ten_ft)).label("fts_count"),
        )
        .group_by(DeviceRecall.cum_xa)
        .order_by(DeviceRecall.cum_xa)
        .all()
    )

    result = []
    for r in rows:
        cum_name = r.cum_xa
        short_code = extract_short_cluster(cum_name)
        tot_dev = int(r.total_devices or 0)
        tot_col = int(r.total_collected or 0)
        result.append({
            "cum_xa": cum_name,
            "short_code": short_code,
            "subscribers_count": int(r.subscribers_count or 0),
            "total_devices": tot_dev,
            "total_collected": tot_col,
            "remaining_devices": max(0, tot_dev - tot_col),
            "fts_count": int(r.fts_count or 0),
        })

    return result


def get_device_recall_fts_by_cluster(db: Session, cluster: str) -> List[dict]:
    """
    Lấy danh sách các nhân viên (Tên FT, MNV) trong cụm được chọn.
    """
    if not cluster:
        return []

    rows = (
        db.query(
            DeviceRecall.ten_ft,
            DeviceRecall.ma_nv,
            func.count(DeviceRecall.id).label("subscribers_count"),
            func.sum(DeviceRecall.tong_tbi_phai_thu).label("total_devices"),
            func.sum(DeviceRecall.sl_thiet_bi_da_thu).label("total_collected"),
        )
        .filter(DeviceRecall.cum_xa == cluster)
        .group_by(DeviceRecall.ten_ft, DeviceRecall.ma_nv)
        .order_by(DeviceRecall.ten_ft)
        .all()
    )

    result = []
    for r in rows:
        tot_dev = int(r.total_devices or 0)
        tot_col = int(r.total_collected or 0)
        result.append({
            "ten_ft": r.ten_ft,
            "ma_nv": r.ma_nv or "",
            "subscribers_count": int(r.subscribers_count or 0),
            "total_devices": tot_dev,
            "total_collected": tot_col,
            "remaining_devices": max(0, tot_dev - tot_col),
        })

    return result


def get_device_recall_items(
    db: Session,
    cluster: str,
    ft_name: Optional[str] = None,
    search: Optional[str] = None
) -> dict:
    """
    Lấy danh sách chi tiết các thiết bị cần thu hồi cho cụm và FT đã chọn.
    Đã lược bỏ Mức TB, Tuổi thọ Mesh, Mesh phải thu theo format mới.
    Bổ sung SL thiết bị đã thu và Điện thoại liên hệ.
    """
    if not cluster:
        return {
            "cluster": "",
            "ft_name": "",
            "items": [],
            "summary": {
                "total_subscribers": 0,
                "total_devices": 0,
                "total_collected": 0,
                "remaining_devices": 0,
                "mesh_devices": 0,
                "regular_devices": 0,
            }
        }

    query = db.query(DeviceRecall).filter(DeviceRecall.cum_xa == cluster)

    if ft_name:
        query = query.filter(DeviceRecall.ten_ft == ft_name)

    if search:
        s = f"%{search.strip()}%"
        query = query.filter(
            (DeviceRecall.so_thue_bao.like(s)) |
            (DeviceRecall.dia_chi_khach_hang.like(s)) |
            (DeviceRecall.dich_vu.like(s)) |
            (DeviceRecall.dien_thoai_lh.like(s))
        )

    items_db = query.order_by(DeviceRecall.tong_tbi_phai_thu.desc(), DeviceRecall.so_thue_bao.asc()).all()

    items = []
    total_devices = 0
    total_collected = 0
    mesh_devices = 0
    regular_devices = 0

    for item in items_db:
        tbi_val = item.tong_tbi_phai_thu or 0
        da_thu_val = item.sl_thiet_bi_da_thu or 0
        total_devices += tbi_val
        total_collected += da_thu_val

        is_mesh = (item.loai_thiet_bi and "mesh" in str(item.loai_thiet_bi).lower())
        if is_mesh:
            mesh_devices += tbi_val
        else:
            regular_devices += tbi_val

        items.append({
            "id": item.id,
            "so_thue_bao": item.so_thue_bao,
            "dich_vu": item.dich_vu or "-",
            "dia_chi_khach_hang": item.dia_chi_khach_hang or "-",
            "dien_thoai_lh": item.dien_thoai_lh or "",
            "user_thanh_ly": item.user_thanh_ly or "",
            "muc_thiet_bi": item.muc_thiet_bi or "-",
            "tong_tbi_phai_thu": tbi_val,
            "sl_thiet_bi_da_thu": da_thu_val,
            "loai_thiet_bi": "MESH" if is_mesh else "Thường (0)",
            "loai_thiet_bi_raw": item.loai_thiet_bi or "-",
            "ma_nv": item.ma_nv or "",
            "ten_ft": item.ten_ft or "",
            "cum_xa": item.cum_xa or "",
        })

    return {
        "cluster": cluster,
        "ft_name": ft_name or "",
        "items": items,
        "summary": {
            "total_subscribers": len(items),
            "total_devices": total_devices,
            "total_collected": total_collected,
            "remaining_devices": max(0, total_devices - total_collected),
            "mesh_devices": mesh_devices,
            "regular_devices": regular_devices,
        }
    }


def export_device_recall_csv(
    db: Session,
    cluster: str,
    ft_name: Optional[str] = None,
    search: Optional[str] = None
) -> str:
    """
    Xuất danh sách thiết bị cần thu hồi dưới dạng CSV (có UTF-8 BOM cho Excel).
    Theo format mới: Bỏ Mức TB, Tuổi thọ Mesh, Mesh phải thu; Thêm SL thiết bị đã thu & Điện thoại LH.
    """
    data = get_device_recall_items(db, cluster=cluster, ft_name=ft_name, search=search)
    items = data.get("items", [])

    output = io.StringIO()
    output.write('\ufeff')
    writer = csv.writer(output)

    writer.writerow([
        "STT",
        "Mã nhân viên",
        "Họ và tên",
        "Cụm xã",
        "Số Thuê Bao",
        "Điện Thoại LH",
        "Dịch Vụ",
        "Địa Chỉ Khách Hàng",
        "Mức Thiết Bị",
        "Tổng Tbi Phải Thu",
        "SL Thiết Bị Đã Thu",
        "Loại Thiết Bị",
    ])

    for idx, it in enumerate(items, 1):
        writer.writerow([
            idx,
            it.get("ma_nv", ""),
            it.get("ten_ft", ""),
            it.get("cum_xa", ""),
            it.get("so_thue_bao", ""),
            it.get("dien_thoai_lh", ""),
            it.get("dich_vu", ""),
            it.get("dia_chi_khach_hang", ""),
            it.get("muc_thiet_bi", ""),
            it.get("tong_tbi_phai_thu", 0),
            it.get("sl_thiet_bi_da_thu", 0),
            it.get("loai_thiet_bi", ""),
        ])

    return output.getvalue()


def export_device_recall_xlsx(
    db: Session,
    cluster: str,
    ft_name: Optional[str] = None,
    search: Optional[str] = None
) -> bytes:
    """
    Xuất danh sách thiết bị cần thu hồi ra file Excel (.xlsx).
    Theo format mới: Bỏ Mức TB, Tuổi thọ Mesh, Mesh phải thu; Thêm SL thiết bị đã thu & Điện thoại LH.
    """
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    data = get_device_recall_items(db, cluster=cluster, ft_name=ft_name, search=search)
    items = data.get("items", [])

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Thu hồi thiết bị"
    ws.views.sheetView[0].showGridLines = True

    headers = [
        "STT",
        "Mã nhân viên",
        "Họ và tên",
        "Cụm xã",
        "Số Thuê Bao",
        "Điện Thoại LH",
        "Dịch Vụ",
        "Địa Chỉ Khách Hàng",
        "Mức Thiết Bị",
        "Tổng Tbi Phải Thu",
        "SL Thiết Bị Đã Thu",
        "Loại Thiết Bị",
    ]
    ws.append(headers)
    ws.row_dimensions[1].height = 26
    ws.freeze_panes = "A2"

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    thin_border_side = Side(border_style="thin", color="CBD5E1")
    cell_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    center_align = Alignment(horizontal="center", vertical="center")

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = cell_border

    for idx, it in enumerate(items, 1):
        ws.append([
            idx,
            it.get("ma_nv", ""),
            it.get("ten_ft", ""),
            it.get("cum_xa", ""),
            it.get("so_thue_bao", ""),
            it.get("dien_thoai_lh", ""),
            it.get("dich_vu", ""),
            it.get("dia_chi_khach_hang", ""),
            it.get("muc_thiet_bi", ""),
            it.get("tong_tbi_phai_thu", 0),
            it.get("sl_thiet_bi_da_thu", 0),
            it.get("loai_thiet_bi", ""),
        ])

    for r_idx in range(2, len(items) + 2):
        ws.row_dimensions[r_idx].height = 20
        for c_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=r_idx, column=c_idx)
            cell.border = cell_border
            cell.font = Font(name="Calibri", size=10)
            if isinstance(cell.value, (int, float)):
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = str(cell.value or "")
            if len(val) > max_len:
                max_len = len(val)
        ws.column_dimensions[col_letter].width = max(10, min(max_len + 4, 60))

    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def get_device_recall_meta(db: Session) -> dict:
    """Lấy thông tin metadata về lần cập nhật gần nhất."""
    total_count = _get_setting(db, "device_recall_total_count", "0")
    total_int = int(total_count) if total_count.isdigit() else 0

    # Auto-seed if empty and file exists (ưu tiên thuhoithietbi_update.txt)
    if total_int == 0:
        seed_file = None
        for fn in ("thuhoithietbi_update.txt", "thuhoithietbi.txt"):
            fp = BASE_DIR / fn
            if fp.exists():
                seed_file = fp
                break
        if seed_file:
            try:
                with open(seed_file, "r", encoding="utf-8", errors="replace") as f:
                    content = f.read()
                import_device_recall_content(db, content, seed_file.name)
                total_int = int(_get_setting(db, "device_recall_total_count", "0"))
            except Exception as e:
                print(f"Auto seed error: {e}")

    return {
        "last_import_time": _get_setting(db, "device_recall_last_import_time"),
        "last_import_time_vn": _get_setting(db, "device_recall_last_import_time_vn"),
        "filename": _get_setting(db, "device_recall_last_filename", "thuhoithietbi_update.txt"),
        "total_count": total_int,
        "clusters_count": int(_get_setting(db, "device_recall_clusters_count", "0")),
        "fts_count": int(_get_setting(db, "device_recall_fts_count", "0")),
        "total_devices": int(_get_setting(db, "device_recall_total_devices", "0")),
        "total_collected": int(_get_setting(db, "device_recall_total_collected", "0")),
        "skipped_count": int(_get_setting(db, "device_recall_skipped_count", "0")),
    }
