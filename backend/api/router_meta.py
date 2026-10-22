from io import BytesIO
from urllib.parse import quote
from typing import List, Any, Optional
from pydantic import BaseModel
from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import distinct
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from backend.database import get_db
from backend.models import Group, Employee, SystemModel, TaskType, Task

router = APIRouter(prefix="/api/meta", tags=["Metadata"])


class ExportSheetData(BaseModel):
    name: Optional[str] = "Sheet1"
    headers: List[str]
    rows: List[List[Any]]


class ExportXlsxPayload(BaseModel):
    filename: Optional[str] = "export"
    sheets: List[ExportSheetData]


@router.post("/export-xlsx")
def export_xlsx_endpoint(payload: ExportXlsxPayload):
    """
    Tạo và tải về file Excel (.xlsx) chuẩn Microsoft Excel, định dạng header chuyên nghiệp,
    tự căn chỉnh độ rộng cột, hiển thị lưới kẻ ô và hỗ trợ font tiếng Việt chuẩn.
    """
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # Xóa sheet mặc định

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    thin_border_side = Side(border_style="thin", color="CBD5E1")
    cell_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=False)
    left_align = Alignment(horizontal="left", vertical="center", wrap_text=False)
    right_align = Alignment(horizontal="right", vertical="center", wrap_text=False)

    for s_idx, sheet_data in enumerate(payload.sheets):
        raw_name = sheet_data.name or f"Sheet{s_idx + 1}"
        for ch in [":", "\\", "/", "?", "*", "[", "]"]:
            raw_name = raw_name.replace(ch, "_")
        sheet_title = raw_name[:31]

        ws = wb.create_sheet(title=sheet_title)
        ws.views.sheetView[0].showGridLines = True

        # Ghi hàng tiêu đề
        ws.append(sheet_data.headers)
        ws.row_dimensions[1].height = 26
        ws.freeze_panes = "A2"

        for col_idx in range(1, len(sheet_data.headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align
            cell.border = cell_border

        # Ghi các hàng dữ liệu
        for row in sheet_data.rows:
            cleaned_row = []
            for item in row:
                if item is None:
                    cleaned_row.append("")
                else:
                    cleaned_row.append(item)
            ws.append(cleaned_row)

        total_rows = len(sheet_data.rows)
        for r_idx in range(2, total_rows + 2):
            ws.row_dimensions[r_idx].height = 20
            for c_idx in range(1, len(sheet_data.headers) + 1):
                cell = ws.cell(row=r_idx, column=c_idx)
                cell.border = cell_border
                cell.font = Font(name="Calibri", size=10)
                if isinstance(cell.value, (int, float)):
                    cell.alignment = right_align
                else:
                    cell.alignment = left_align

        # Tự tính độ rộng cột
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or "")
                first_line = val.split("\n")[0]
                if len(first_line) > max_len:
                    max_len = len(first_line)
            ws.column_dimensions[col_letter].width = max(10, min(max_len + 4, 60))

    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)

    clean_fn = payload.filename or "export"
    if clean_fn.endswith(".xlsx"):
        clean_fn = clean_fn[:-5]
    elif clean_fn.endswith(".csv"):
        clean_fn = clean_fn[:-4]
    final_filename = f"{clean_fn}.xlsx"

    return StreamingResponse(
        bio,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f"attachment; filename*=UTF-8''{quote(final_filename)}",
            "Access-Control-Expose-Headers": "Content-Disposition",
        }
    )


@router.get("/filters")
def get_filter_options(db: Session = Depends(get_db)):
    """Fetch distinct lists for select dropdowns in filters."""
    groups = db.query(Group.id, Group.name).order_by(Group.name).all()
    employees = db.query(Employee.id, Employee.name).order_by(Employee.name).all()
    systems = db.query(SystemModel.id, SystemModel.name).order_by(SystemModel.name).all()

    # Get distinct task types directly from tasks table for accuracy
    task_types_q = db.query(distinct(Task.loai_cong_viec))\
        .filter(Task.loai_cong_viec != None)\
        .order_by(Task.loai_cong_viec).all()
    task_types = [t[0] for t in task_types_q if t[0]]

    # Get distinct statuses
    statuses_q = db.query(distinct(Task.trang_thai))\
        .filter(Task.trang_thai != None)\
        .order_by(Task.trang_thai).all()
    statuses = [s[0] for s in statuses_q if s[0]]

    return {
        "groups": [{"id": g.id, "name": g.name} for g in groups],
        "employees": [{"id": e.id, "name": e.name} for e in employees],
        "systems": [{"id": s.id, "name": s.name} for s in systems],
        "task_types": task_types,
        "statuses": statuses,
    }
